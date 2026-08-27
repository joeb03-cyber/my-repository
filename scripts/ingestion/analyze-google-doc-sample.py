#!/usr/bin/env python3
"""Summarize parser-relevant structure without copying document text to output."""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document


if len(sys.argv) != 3:
    raise SystemExit("Usage: analyze-google-doc-sample.py <sample-dir> <output.json>")

sample_dir = Path(sys.argv[1])
output_path = Path(sys.argv[2])
manifest = json.loads((sample_dir / "sample-download-manifest.json").read_text())

location_pattern = re.compile(r"\b(?:page|p\.|location|loc\.)\s*[:#-]?\s*\d+\b", re.I)
separator_pattern = re.compile(r"^(?:[-_*•·—–=]\s*){3,}$")


def has_numbering(paragraph) -> bool:
    properties = paragraph._p.pPr
    return properties is not None and properties.numPr is not None


documents = []
for item in manifest["results"]:
    if not item.get("downloaded"):
        continue

    doc = Document(sample_dir / item["local_file"])
    nonempty = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip()]
    styles = Counter(paragraph.style.name for paragraph in nonempty)
    heading_paragraphs = [paragraph for paragraph in nonempty if paragraph.style.name.lower().startswith("heading")]
    numbered = [paragraph for paragraph in nonempty if has_numbering(paragraph)]
    location_markers = [paragraph for paragraph in nonempty if location_pattern.search(paragraph.text)]
    separators = [paragraph for paragraph in nonempty if separator_pattern.fullmatch(paragraph.text.strip())]
    quote_led = [paragraph for paragraph in nonempty if paragraph.text.lstrip().startswith(('"', "“", "‘", "'"))]
    paragraphs_with_bold = sum(any(run.bold for run in paragraph.runs) for paragraph in nonempty)
    paragraphs_with_italic = sum(any(run.italic for run in paragraph.runs) for paragraph in nonempty)
    paragraphs_with_highlight = sum(
        any(run.font.highlight_color is not None for run in paragraph.runs) for paragraph in nonempty
    )
    paragraphs_with_mixed_formatting = sum(
        len({(bool(run.bold), bool(run.italic), run.font.highlight_color is not None) for run in paragraph.runs if run.text}) > 1
        for paragraph in nonempty
    )

    documents.append(
        {
            **{key: value for key, value in item.items() if key != "local_file"},
            "sample_file": item["local_file"],
            "paragraphs_total": len(doc.paragraphs),
            "paragraphs_nonempty": len(nonempty),
            "word_count_approx": sum(len(paragraph.text.split()) for paragraph in nonempty),
            "paragraph_style_counts": dict(styles),
            "heading_paragraphs": len(heading_paragraphs),
            "numbered_or_bulleted_paragraphs": len(numbered),
            "table_count": len(doc.tables),
            "location_marker_paragraphs": len(location_markers),
            "separator_paragraphs": len(separators),
            "quote_led_paragraphs": len(quote_led),
            "paragraphs_with_bold": paragraphs_with_bold,
            "paragraphs_with_italic": paragraphs_with_italic,
            "paragraphs_with_text_highlight": paragraphs_with_highlight,
            "paragraphs_with_mixed_run_formatting": paragraphs_with_mixed_formatting,
            "section_count": len(doc.sections),
        }
    )

output = {
    "schema_version": "bookshelf-google-doc-sample-analysis.v1",
    "method": "python-docx structural analysis; document text is not copied into this output",
    "sample_count": len(documents),
    "documents": documents,
}
output_path.write_text(json.dumps(output, indent=2) + "\n")
print(json.dumps(output, indent=2))
